"""Create an instructor-only Excel snapshot of persisted assessment scores."""

from datetime import datetime, timezone
from io import BytesIO
from math import isfinite

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


NAVY = "1F2874"
YELLOW = "FFE833"
TEXT = "1F2430"
MUTED = "667085"
STRIPE = "F5F6FA"


def _safe_text(value):
    """Keep untrusted database text from becoming an Excel formula."""
    text = str(value or "")
    if text.lstrip().startswith(("=", "+", "-", "@")) or text[:1] in ("\t", "\r", "\n"):
        return "'" + text
    return text


def _percentage(grade):
    stored = grade.get("percentage")
    if stored is not None:
        try:
            value = float(stored) / 100
            return value if isfinite(value) else None
        except (TypeError, ValueError):
            pass
    try:
        score = float(grade["score"])
        total = float(grade["total_questions"])
        value = score / total if total > 0 else None
        return value if value is not None and isfinite(value) else None
    except (KeyError, TypeError, ValueError):
        return None


def _utc_datetime(value):
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone(timezone.utc).replace(tzinfo=None)
    except ValueError:
        return _safe_text(value)


def build_scores_workbook(assessment_title, class_name, submissions,
                          results_released=False, exported_at=None):
    """Return a seeked XLSX buffer with one row per saved graded sheet."""
    exported_at = exported_at or datetime.now(timezone.utc)
    if exported_at.tzinfo is not None:
        exported_at = exported_at.astimezone(timezone.utc).replace(tzinfo=None)
    rows = sorted(submissions, key=lambda row: (
        str(row.get("student_name") or "").casefold(),
        str(row.get("student_email") or "").casefold(),
        str(row.get("sheet_code") or ""),
    ))
    percentages = [_percentage(row) for row in rows]
    valid_percentages = [value for value in percentages if value is not None]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Scores"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = NAVY
    sheet.freeze_panes = "C10"

    widths = [26, 31, 9, 10, 16, 23, 22, 20]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet["A2"] = "CheckMate LMS · Assessment scores"
    sheet["A2"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
    sheet.row_dimensions[2].height = 25
    for cell in sheet[2][:8]:
        cell.border = Border(bottom=Side(style="thin", color="D7DAE2"))

    metadata = {
        "A3": "Assessment", "B3": _safe_text(assessment_title),
        "A4": "Class", "B4": _safe_text(class_name),
        "E3": "Exported (UTC)", "F3": exported_at,
        "E4": "Results", "F4": "Released" if results_released else "Instructor review",
    }
    for address, value in metadata.items():
        cell = sheet[address]
        cell.value = value
        cell.font = Font(name="Arial", size=10, color=MUTED if address in ("A3", "A4", "E3", "E4") else TEXT)
        cell.alignment = Alignment(vertical="center")
    sheet["F3"].number_format = 'yyyy-mm-dd hh:mm "UTC"'

    summary = [
        ("A6", "Graded sheets", "B6", len(rows)),
        ("C6", "Average", "D6", sum(valid_percentages) / len(valid_percentages) if valid_percentages else None),
        ("E6", "Highest", "F6", max(valid_percentages) if valid_percentages else None),
        ("G6", "Lowest", "H6", min(valid_percentages) if valid_percentages else None),
    ]
    for label_address, label, value_address, value in summary:
        label_cell = sheet[label_address]
        label_cell.value = label
        label_cell.font = Font(name="Arial", size=10, color=MUTED)
        value_cell = sheet[value_address]
        value_cell.value = value
        value_cell.font = Font(name="Arial", size=11, bold=True, color=NAVY)
        value_cell.alignment = Alignment(horizontal="right", vertical="center")
        if value_address != "B6":
            value_cell.number_format = "0.0%"
    sheet.row_dimensions[6].height = 23
    sheet["A7"] = "Saved graded sheets only · Export again to include later scans."
    sheet["A7"].font = Font(name="Arial", size=9, italic=True, color=MUTED)

    headers = ["Student", "Email", "Set", "Score", "Out of", "Percent", "Graded (UTC)", "Sheet code"]
    for column, label in enumerate(headers, 1):
        cell = sheet.cell(9, column, label)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=YELLOW))
    sheet.row_dimensions[9].height = 25

    for offset, (row, percentage) in enumerate(zip(rows, percentages), 10):
        values = [
            _safe_text(row.get("student_name") or "Student"),
            _safe_text(row.get("student_email")),
            _safe_text(row.get("set_type") or "A"),
            row.get("score"),
            row.get("total_questions"),
            percentage,
            _utc_datetime(row.get("created_at")),
            _safe_text(row.get("sheet_code")),
        ]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(offset, column, value)
            cell.font = Font(name="Arial", size=10, color=TEXT)
            cell.alignment = Alignment(
                horizontal="right" if column in (4, 5, 6) else "left",
                vertical="center",
            )
            if offset % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=STRIPE)
            if column == 6:
                cell.number_format = "0.0%"
            elif column == 7 and isinstance(value, datetime):
                cell.number_format = 'yyyy-mm-dd hh:mm "UTC"'
            elif column in (4, 5):
                cell.number_format = "0"
        sheet.row_dimensions[offset].height = 21

    sheet.auto_filter.ref = f"A9:H{max(9, 9 + len(rows))}"
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:9"
    sheet.print_area = f"A2:H{max(9, 9 + len(rows))}"

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
