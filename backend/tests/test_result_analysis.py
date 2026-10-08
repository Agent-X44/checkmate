import copy
import json
import os
import sys
from datetime import datetime
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from pydantic import ValidationError

# Tests are entirely offline, including module initialization.
os.environ["SUPABASE_URL"] = ""
os.environ["SUPABASE_KEY"] = ""
os.environ["OPENROUTER_API_KEY"] = ""
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import main
import ai_service


class Query:
    def __init__(self, db, table):
        self.db, self.table = db, table
        self.filters, self.action, self.payload = [], "select", None
        self.columns, self.sort, self.bounds = "*", None, None

    def select(self, columns="*"):
        self.columns = columns
        return self

    def eq(self, key, value):
        if key == 'answers' and isinstance(value, str): value = json.loads(value)
        self.filters.append((key, value))
        return self

    def order(self, key, desc=False):
        self.sort = key, desc
        return self

    def range(self, start, end):
        self.bounds = start, end + 1
        return self

    def limit(self, count):
        self.bounds = 0, count
        return self

    def update(self, payload):
        self.action, self.payload = "update", payload
        return self

    def insert(self, payload):
        self.action, self.payload = "insert", payload
        return self

    def execute(self):
        self.db.calls.append((self.table, self.action, self.filters, self.payload))
        if self.table == self.db.fail_table:
            raise RuntimeError("Database failure")
        rows = [row for row in self.db.rows[self.table]
                if all(row.get(key) == value for key, value in self.filters)]
        if self.action == "insert":
            row = {"id": "saved-grade", **copy.deepcopy(self.payload)}
            self.db.rows[self.table].append(row)
            rows = [row]
        elif self.action == "update":
            for row in rows:
                row.update(copy.deepcopy(self.payload))
        if self.sort:
            key, desc = self.sort
            rows = sorted(rows, key=lambda row: row.get(key) or "", reverse=desc)
        if self.bounds:
            rows = rows[self.bounds[0]:self.bounds[1]]
        rows = copy.deepcopy(rows)
        if self.table == "answer_sheets" and "exams(" in self.columns:
            for row in rows:
                row.setdefault("exams", next((exam for exam in self.db.rows["exams"]
                                             if exam["id"] == row.get("exam_id")), None))
        if self.table == "answer_sheets" and "grades(" in self.columns:
            for row in rows:
                grades = [grade for grade in self.db.rows["grades"] if grade["sheet_id"] == row["id"]]
                row["grades"] = grades[0] if self.db.single_relation and len(grades) == 1 else grades
        return SimpleNamespace(data=rows)


class Database:
    def __init__(self):
        self.calls = []
        self.fail_table = None
        self.single_relation = False
        self.rows = {
            "classes": [{"id": "class-1", "name": "Geography 3A", "instructor_id": "teacher"}],
            "enrollments": [{"id": "enrollment-1", "class_id": "class-1", "user_id": "student-1"}],
            "questions": [
                {"id": "q-1", "exam_id": "exam-1", "question_text": "Capital of France?",
                 "question_type": "MCQ", "correct_answer": "A", "topic_tag": "Capitals"},
                {"id": "q-2", "exam_id": "exam-1", "question_text": "Largest ocean?",
                 "question_type": "MCQ", "correct_answer": "C", "topic_tag": "Oceans"},
            ],
            "exams": [{"id": "exam-1", "class_id": "class-1", "title": "Geography", "results_released": False,
                       "is_approved": True,
                       "classes": {"instructor_id": "teacher"}}],
            "answer_sheets": [{"id": "sheet-1", "sheet_identifier": "CM-ABC12345",
                               "exam_id": "exam-1", "student_id": "student-1",
                               "set_type": "B", "profiles": {"name": "Ada Student", "email": "ada@example.test"}}],
            "grades": [{"id": "grade-1", "sheet_id": "sheet-1", "score": 1,
                        "total_questions": 2, "created_at": "2026-09-26", "answers": [
                            {"question_number": 1, "question_id": "q-1", "question_text": "Capital of France?",
                             "topic_tag": "Capitals", "answer": "A", "correct_answer": "A", "isCorrect": True},
                            {"question_number": 2, "question_id": "q-2", "question_text": "Largest ocean?",
                             "topic_tag": "Oceans", "answer": "B", "correct_answer": "C", "isCorrect": False},
                        ]}],
        }

    def table(self, name):
        return Query(self, name)

    def rpc(self, name, params):
        assert name == "save_grade_session"

        def execute():
            self.calls.append(("grades", "rpc", [], params))
            if self.fail_table == "grades":
                raise RuntimeError("Database failure")
            for payload in params["p_results"]:
                matches = [row for row in self.rows["grades"] if row["sheet_id"] == payload["sheet_id"]]
                if not matches:
                    matches = [{"id": "saved-grade"}]
                    self.rows["grades"].extend(matches)
                for row in matches:
                    row.update(copy.deepcopy(payload), student_insight=None)
            return SimpleNamespace(data=len(params["p_results"]))

        return SimpleNamespace(execute=execute)


@pytest.fixture
def setup(monkeypatch):
    db = Database()
    monkeypatch.setattr(main, "supabase", db)
    user = SimpleNamespace(user=SimpleNamespace(id="teacher"))
    main.app.dependency_overrides[main.get_current_user] = lambda: user
    ai = AsyncMock(return_value=main.ClassAnalysisResponse(
        insights="One submission: review the missed ocean item.",
        recommendations="Practice identifying oceans on a map.",
    ))
    monkeypatch.setattr(main, "generate_structured_response", ai)
    yield db, user, ai
    main.app.dependency_overrides.clear()


async def post(path, payload):
    async with AsyncClient(transport=ASGITransport(app=main.app), base_url="http://test") as client:
        return await client.post(path, json=payload)


async def get(path):
    async with AsyncClient(transport=ASGITransport(app=main.app), base_url="http://test") as client:
        return await client.get(path)


@pytest.mark.asyncio
async def test_student_result_requires_release_and_returns_own_saved_items(setup):
    db, user, _ = setup
    user.user.id = "student-1"
    assert (await get("/my-exam-result/exam-1")).status_code == 404
    db.rows["exams"][0]["results_released"] = True
    result = await get("/my-exam-result/exam-1")
    assert result.status_code == 200
    assert result.json()["grade"]["sheet_id"] == "sheet-1"
    assert result.json()["grade"]["answers"][1]["answer"] == "B"
    assert result.json()["grade"]["answers"][1]["correct_answer"] == "C"
    assert result.json()["grade"]["answers"][1]["question_text"] == "Largest ocean?"
    user.user.id = "student-2"
    assert (await get("/my-exam-result/exam-1")).status_code == 404


def test_released_item_detail_preserves_ambiguous_selected_marks():
    from result_analysis import enrich_answers

    grade = {"answers": [{"question_id": "q-1", "answer": None,
                          "isAmbiguous": True, "multipleAnswers": ["A", "C"]}]}
    answer = enrich_answers(grade, {"q-1": {"correct_answer": "A"}})["answers"][0]
    assert answer["multipleAnswers"] == ["A", "C"]
    assert answer["correct_answer"] == "A"


@pytest.mark.asyncio
async def test_student_exam_list_excludes_drafts_and_other_users(setup):
    db, user, _ = setup
    db.rows["exams"].append({"id": "exam-draft", "class_id": "class-1",
                             "title": "Draft", "is_approved": False})
    user.user.id = "teacher"
    assert len((await get("/get-exams/class-1")).json()) == 2
    user.user.id = "student-1"
    response = await get("/get-exams/class-1")
    assert response.status_code == 200
    assert [exam["id"] for exam in response.json()] == ["exam-1"]
    user.user.id = "student-2"
    assert (await get("/get-exams/class-1")).status_code == 403


@pytest.mark.asyncio
async def test_instructor_exports_formatted_saved_scores(setup):
    db, _, _ = setup
    db.rows["grades"].append({"id": "older", "sheet_id": "sheet-1",
                              "score": 0, "total_questions": 2,
                              "created_at": "2026-09-25"})
    response = await get("/export-scores/exam-1")
    assert response.status_code == 200
    assert "spreadsheetml.sheet" in response.headers["content-type"]
    assert response.content[:2] == b"PK"
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook["Scores"]
    assert sheet["B3"].value == "Geography"
    assert sheet["B4"].value == "Geography 3A"
    assert sheet["B6"].value == 1
    assert sheet["D6"].value == 0.5
    assert [sheet.cell(10, col).value for col in range(1, 9)] == [
        "Ada Student", "ada@example.test", "B", 1, 2, 0.5,
        datetime(2026, 9, 26), "CM-ABC12345",
    ]
    assert sheet["F10"].number_format == "0.0%"
    assert sheet["A9"].fill.fgColor.rgb.endswith("1F2874")
    assert sheet.auto_filter.ref == "A9:H10"
    assert sheet.freeze_panes == "C10"
    assert sheet.sheet_view.showGridLines is False


@pytest.mark.asyncio
async def test_student_cannot_export_class_scores(setup):
    _, user, _ = setup
    user.user.id = "student-1"
    assert (await get("/export-scores/exam-1")).status_code == 403


def test_spreadsheet_text_cannot_become_formula():
    from scores_export import build_scores_workbook

    data = build_scores_workbook("=Hidden formula", "Class", [{
        "student_name": "=HYPERLINK(\"https://example.test\")",
        "student_email": "student@example.test", "score": 0,
        "total_questions": 1, "percentage": 0,
    }])
    sheet = load_workbook(data).active
    assert sheet["B3"].data_type == "s"
    assert sheet["A10"].data_type == "s"
    assert sheet["D10"].value == 0
    assert sheet["F10"].value == 0


@pytest.mark.asyncio
@pytest.mark.parametrize("single_relation", [False, True])
async def test_one_result_always_invokes_ai_with_actual_item_context(setup, single_relation):
    db, _, ai = setup
    db.single_relation = single_relation
    response = await post("/analyze-class", {"exam_id": "exam-1"})
    assert response.status_code == 200
    assert response.json()["sample_count"] == 1
    assert response.json()["source"] == "ai"
    assert response.json()["metrics"]["topics"]["Oceans"]["correct"] == 0
    assert any(question["question_text"] == "Largest ocean?"
               for question in response.json()["metrics"]["questions"])
    assert response.json()["analysis"]["insights"]
    ai.assert_awaited_once()
    prompt = ai.call_args.kwargs["user_prompt"]
    assert '"sample_count": 1' in prompt
    assert "Largest ocean?" in prompt and '"isCorrect": false' in prompt


@pytest.mark.asyncio
async def test_analysis_uses_question_table_instead_of_stale_item_snapshot(setup):
    db, _, ai = setup
    db.rows["grades"][0]["answers"][1]["question_text"] = "Stale paper text"
    db.rows["grades"][0]["answers"][1]["topic_tag"] = "Wrong topic"
    response = await post("/analyze-class", {"exam_id": "exam-1"})
    assert response.status_code == 200
    prompt = ai.call_args.kwargs["user_prompt"]
    assert "Largest ocean?" in prompt
    assert "Stale paper text" not in prompt
    assert '"Oceans"' in prompt


@pytest.mark.asyncio
async def test_instructor_can_review_student_course_history_before_release(setup):
    db, _, ai = setup
    response = await post("/student-overall-analysis?student_id=student-1&class_id=class-1", {})
    assert response.status_code == 200
    assert response.json()["sample_count"] == 1
    assert ai.await_count == 1


@pytest.mark.asyncio
async def test_unrelated_user_cannot_review_student_course_history(setup):
    _, user, ai = setup
    user.user.id = "student-2"
    response = await post("/student-overall-analysis?student_id=student-1&class_id=class-1", {})
    assert response.status_code == 403
    ai.assert_not_awaited()


@pytest.mark.asyncio
async def test_empty_assessment_does_not_invent_analysis(setup):
    db, _, ai = setup
    db.rows["grades"] = []
    response = await post("/analyze-class", {"exam_id": "exam-1"})
    assert response.json()["status"] == "no_results"
    ai.assert_not_awaited()


@pytest.mark.asyncio
async def test_class_analysis_paginates_past_500_sheets(setup):
    db, _, _ = setup
    db.rows["answer_sheets"] = [
        {"id": f"sheet-{i:04}", "exam_id": "exam-1"} for i in range(501)
    ]
    db.rows["grades"] = [
        {"sheet_id": row["id"], "score": 1, "total_questions": 1}
        for row in db.rows["answer_sheets"]
    ]
    response = await post("/analyze-class", {"exam_id": "exam-1"})
    assert response.json()["sample_count"] == 501


@pytest.mark.asyncio
async def test_legacy_duplicate_grade_rows_count_as_one_sheet(setup):
    db, _, _ = setup
    older = copy.deepcopy(db.rows["grades"][0])
    older["id"] = "older-grade"
    older["created_at"] = "2026-09-25"
    older["score"] = 0
    db.rows["grades"].append(older)
    response = await post("/analyze-class", {"exam_id": "exam-1"})
    assert response.status_code == 200
    assert response.json()["sample_count"] == 1
    assert response.json()["metrics"]["average_percentage"] == 50.0


@pytest.mark.asyncio
async def test_ai_failure_returns_honest_useful_single_result_summary(setup):
    _, _, ai = setup
    ai.side_effect = RuntimeError("AI offline")
    response = await post("/analyze-class", {"exam_id": "exam-1"})
    data = response.json()
    assert response.status_code == 200
    assert data["source"] == "summary"
    assert "50.0%" in data["analysis"]["insights"]
    assert "one submission" in data["analysis"]["insights"]
    assert "Oceans" in data["analysis"]["recommendations"]


@pytest.mark.asyncio
async def test_legacy_total_only_analysis_does_not_invent_topics(setup):
    db, _, ai = setup
    db.rows["grades"][0].pop("answers")
    ai.side_effect = RuntimeError("AI offline")
    data = (await post("/analyze-class", {"exam_id": "exam-1"})).json()
    assert "Only total scores" in data["analysis"]["insights"]


@pytest.mark.asyncio
async def test_instructor_personal_feedback_uses_selected_grade_and_persists(setup):
    db, _, ai = setup
    ai.return_value = main.StudentInsightResponse(
        performanceSummary="You got the capital question right.", strengths=["Capitals"],
        learningGaps=["Oceans"], actionableSteps=["Review an ocean map."],
    )
    response = await post("/student-insight", {"exam_id": "exam-1", "sheet_id": "sheet-1"})
    assert response.status_code == 200
    assert '"score": 1' in ai.call_args.kwargs["user_prompt"]
    assert "Largest ocean?" in ai.call_args.kwargs["user_prompt"]
    assert db.rows["grades"][0]["student_insight"] == response.json()
    assert db.rows["grades"][0]["score"] == 1


@pytest.mark.asyncio
@pytest.mark.parametrize("user_id,released,expected", [
    ("student-1", False, 403), ("student-2", True, 403),
    ("other-teacher", True, 403), ("student-1", True, 200),
])
async def test_personal_insight_requires_own_released_result(setup, user_id, released, expected):
    db, user, ai = setup
    user.user.id = user_id
    db.rows["exams"][0]["results_released"] = released
    ai.side_effect = RuntimeError("AI offline")
    response = await post("/student-insight", {"exam_id": "exam-1", "sheet_id": "sheet-1"})
    assert response.status_code == expected
    if expected == 403:
        ai.assert_not_awaited()
    else:
        assert response.json()["source"] == "summary"
        assert response.json()["insight"]["actionableSteps"]


@pytest.mark.asyncio
async def test_cached_ai_insight_is_reused(setup):
    db, _, ai = setup
    cached = {"source": "ai", "insight": {"performanceSummary": "Saved feedback"}}
    db.rows["grades"][0]["student_insight"] = cached
    response = await post("/student-insight", {"exam_id": "exam-1", "sheet_id": "sheet-1"})
    assert response.json() == cached
    ai.assert_not_awaited()


@pytest.mark.asyncio
async def test_student_overall_analysis_aggregates_released_results(setup):
    db, user, ai = setup
    user.user.id = "student-1"
    db.rows["exams"][0]["results_released"] = True
    db.rows["answer_sheets"].append({
        "id": "sheet-2", "sheet_identifier": "CM-DEF45678",
        "exam_id": "exam-1", "student_id": "student-1",
    })
    db.rows["grades"].append({
        "id": "grade-2", "sheet_id": "sheet-2", "score": 2,
        "total_questions": 2, "created_at": "2026-09-27", "answers": [
            {"question_number": 1, "topic_tag": "Capitals", "answer": "A", "correct_answer": "A", "isCorrect": True},
            {"question_number": 2, "topic_tag": "Oceans", "answer": "C", "correct_answer": "C", "isCorrect": True},
        ],
    })
    ai.return_value = main.StudentOverviewResponse(
        performanceSummary="Across two assessments you averaged 75%.",
        strengths=["Capitals"], learningGaps=["Oceans"],
        actionableSteps=["Review oceans."],
    )
    response = await post("/student-overall-analysis", {})
    data = response.json()
    assert response.status_code == 200
    assert data["sample_count"] == 2
    assert data["source"] == "ai"
    assert data["metrics"]["topics"]["Oceans"]["total"] == 2
    ai.assert_awaited_once()
    assert '"sample_count": 2' in ai.call_args.kwargs["user_prompt"]


@pytest.mark.asyncio
async def test_student_overall_analysis_no_released_results(setup):
    db, user, ai = setup
    user.user.id = "student-1"
    response = await post("/student-overall-analysis", {})
    assert response.json()["status"] == "no_results"
    ai.assert_not_awaited()


@pytest.mark.asyncio
async def test_student_overall_analysis_summary_fallback(setup):
    db, user, ai = setup
    user.user.id = "student-1"
    db.rows["exams"][0]["results_released"] = True
    ai.side_effect = RuntimeError("AI offline")
    response = await post("/student-overall-analysis", {})
    data = response.json()
    assert data["source"] == "summary"
    assert "1 released assessment" in data["analysis"]["performanceSummary"]
    assert "Oceans" in data["analysis"]["learningGaps"][0]


@pytest.mark.asyncio
@pytest.mark.parametrize("path,payload", [
    ("/analyze-class", {"exam_id": "exam-1"}),
    ("/release-results/exam-1", {}),
    ("/batch-save-grades", {"exam_id": "exam-1", "results": [{"sheet_id": "CM-ABC12345", "score": 1, "total": 2}]}),
])
async def test_student_cannot_perform_instructor_actions(setup, path, payload):
    _, user, ai = setup
    user.user.id = "student-1"
    assert (await post(path, payload)).status_code == 403
    ai.assert_not_awaited()


@pytest.mark.asyncio
async def test_anonymous_analysis_is_rejected(setup):
    main.app.dependency_overrides.clear()
    assert (await post("/analyze-class", {"exam_id": "exam-1"})).status_code == 401


@pytest.mark.asyncio
async def test_sync_preserves_local_evaluation_without_ai_or_raw_images(setup):
    db, _, ai = setup
    answers = copy.deepcopy(db.rows["grades"][0]["answers"])
    answers[0]["warpedImage"] = "must-never-be-persisted"
    response = await post("/batch-save-grades", {"exam_id": "exam-1", "results": [
        {"sheet_id": "CM-ABC12345", "score": 1, "total": 2, "answers": answers},
    ]})
    assert response.json() == {"status": "success", "saved_count": 1}
    saved = db.rows["grades"][0]
    assert saved["answers"][1]["answer"] == "B"
    assert saved["answers"][1]["correct_answer"] == "C"
    assert saved["answers"][1]["isCorrect"] is False
    assert "warpedImage" not in json.dumps(saved)
    assert saved["student_insight"] is None
    ai.assert_not_awaited()


@pytest.mark.asyncio
async def test_independent_sheet_sync_accepts_mixed_assessments_and_sets(setup):
    db, _, ai = setup
    db.rows["exams"].append({"id": "exam-2", "class_id": "class-1", "title": "Math",
                             "results_released": False,
                             "classes": {"instructor_id": "teacher"}})
    db.rows["answer_sheets"].append({"id": "sheet-2", "sheet_identifier": "CM-DEF45678",
                                     "exam_id": "exam-2", "student_id": "student-1", "set_type": "B"})
    second = await post("/batch-save-grades", {"exam_id": "exam-2", "results": [
        {"sheet_id": "CM-DEF45678", "score": 1, "total": 1,
         "answers": [{"question_id": "q-3", "answer": "B", "isCorrect": True}]},
    ]})
    first = await post("/batch-save-grades", {"exam_id": "exam-1", "results": [
        {"sheet_id": "CM-ABC12345", "score": 1, "total": 2},
    ]})
    assert second.status_code == first.status_code == 200
    assert {grade["sheet_id"] for grade in db.rows["grades"]} == {"sheet-1", "sheet-2"}
    ai.assert_not_awaited()


@pytest.mark.asyncio
async def test_invalid_sheet_prevents_all_session_writes(setup):
    db, _, ai = setup
    response = await post("/batch-save-grades", {"exam_id": "exam-1", "results": [
        {"sheet_id": "CM-ABC12345", "score": 1, "total": 2},
        {"sheet_id": "CM-NOTFOUND", "score": 1, "total": 2},
    ]})
    assert response.status_code == 404
    assert all(call[1] == "select" for call in db.calls)
    ai.assert_not_awaited()


@pytest.mark.asyncio
async def test_sync_database_failure_is_never_reported_as_success(setup):
    db, _, _ = setup
    db.fail_table = "grades"
    response = await post("/batch-save-grades", {"exam_id": "exam-1", "results": [
        {"sheet_id": "CM-ABC12345", "score": 1, "total": 2},
    ]})
    assert response.status_code == 500


@pytest.mark.asyncio
async def test_mismatched_local_score_is_rejected(setup):
    db, _, _ = setup
    response = await post("/batch-save-grades", {"exam_id": "exam-1", "results": [
        {"sheet_id": "CM-ABC12345", "score": 2, "total": 2, "answers": db.rows["grades"][0]["answers"]},
    ]})
    assert response.status_code == 422


@pytest.mark.asyncio
async def test_ai_prompt_contract_matches_actual_validator(monkeypatch):
    payload = {"insights": "One saved result: 50%.", "recommendations": "Review oceans."}
    create = AsyncMock(return_value=SimpleNamespace(choices=[
        SimpleNamespace(message=SimpleNamespace(content=json.dumps(payload)))
    ]))
    monkeypatch.setattr(ai_service, "client", SimpleNamespace(chat=SimpleNamespace(
        completions=SimpleNamespace(create=create))))
    result = await ai_service.generate_structured_response(
        "Analyze saved results", "one result", ai_service.ClassAnalysisResponse)
    assert result.insights == payload["insights"]
    prompt = create.call_args.kwargs["messages"][0]["content"]
    assert '"required": ["insights", "recommendations"]' in prompt
    assert "topicBreakdown" not in prompt
    with pytest.raises(ValidationError):
        ai_service.ClassAnalysisResponse(insights="  ", recommendations="Review")
